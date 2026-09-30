# -*- coding: utf-8 -*-
"""
Created on Tue Sep 29 15:47:21 2026

@author: rmili
"""

import json
import re
from pathlib import Path

import openpyxl


def parse_mapping_comment(comment_text):
    """
    Convert a MAS cell comment into structured mapping metadata.

    Example comment:
        Metric = B1_Amount
        Data Type = Bk_Number 19 - 9
        Bk_Sfp_Assets = CashBalances
    """

    result = {
        "metric": None,
        "datatype": None,
        "dimensions": {}
    }

    if not comment_text:
        return result

    for line in comment_text.splitlines():
        line = line.strip()

        if not line or "=" not in line:
            continue

        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip()

        if key.lower() == "metric":
            result["metric"] = value

        elif key.lower() in ("data type", "datatype"):
            result["datatype"] = value

        else:
            # Anything else expressed as key=value
            # is treated as a dimension.
            result["dimensions"][key] = value

    return result


def extract_mapping(
    input_file,
    sheet_names,
    output_file,
    cell_range=None
):
    """
    Extract MAS XML mappings from comments in selected Excel sheets.

    Parameters
    ----------
    input_file : str
        MAS XML Schema - Excel Submission Template Mapping workbook.

    sheet_names : list[str]
        Sheets to process, e.g. ["B1(1)", "B2(1)"].

    output_file : str
        JSON file to create.

    cell_range : str | None
        Optional Excel range, e.g. "C8:C40".
        If None, all cells in each selected sheet are checked.
    """

    input_path = Path(input_file)

    if not input_path.exists():
        raise FileNotFoundError(
            f"Mapping workbook not found: {input_file}"
        )

    workbook = openpyxl.load_workbook(
        input_file,
        data_only=False
    )

    mappings = []

    for sheet_name in sheet_names:

        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[sheet_name]

        # Parameterised cell selection
        if cell_range:
            cells = (
                cell
                for row in worksheet[cell_range]
                for cell in row
            )
        else:
            cells = (
                cell
                for row in worksheet.iter_rows()
                for cell in row
            )

        for cell in cells:

            # Ignore cells without MAS mapping comments
            if cell.comment is None:
                continue

            parsed = parse_mapping_comment(
                cell.comment.text
            )

            # Ignore comments that aren't mapping comments
            if not parsed["metric"]:
                continue

            mapping = {
                "sheet": sheet_name,
                "cell": cell.coordinate,
                "metric": parsed["metric"],
                "datatype": parsed["datatype"],
                "dimensions": parsed["dimensions"]
            }

            mappings.append(mapping)

    # Write structured mapping
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(
            mappings,
            f,
            indent=4,
            ensure_ascii=False
        )

    print(f"Extracted {len(mappings)} mappings.")
    print(f"Mapping information written to: {output_file}")

    return mappings


